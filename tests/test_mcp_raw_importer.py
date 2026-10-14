"""
Unit & Security PoC Test Suite for MCP import_document_to_markdown tool.
Verifies all safety boundaries, zip bomb defense, session asset containment,
and version history lifecycle.
"""
import os
import sys
import json
import uuid
import docx
import fitz
import shutil
import openpyxl
import tempfile
import unittest
from unittest.mock import patch, MagicMock
from PIL import Image

from src.services.metadata_index import MetadataIndex
from src.services.media_asset_manager import MediaAssetManager
from src.mcp.tools import (
    handle_import_document_to_markdown,
    handle_patch_document_content,
    handle_rollback_document,
    handle_read_document,
    _sanitize_filename_stem,
    _sanitize_error_message
)
from src.mcp.security import is_strictly_within


class TestMCPRawImporter(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.mkdtemp()
        self.workspace_dir = os.path.join(self.tmp_dir, "workspace")
        os.makedirs(self.workspace_dir, exist_ok=True)
        
        self.db_path = os.path.join(self.tmp_dir, "test_index.db")
        self.index = MetadataIndex(self.db_path)
        
        # Set workspace environment for MCP security checks
        self.orig_env = os.environ.get("DOCCONVERT_WORKSPACE")
        os.environ["DOCCONVERT_WORKSPACE"] = self.workspace_dir

    def tearDown(self):
        if self.orig_env is not None:
            os.environ["DOCCONVERT_WORKSPACE"] = self.orig_env
        else:
            os.environ.pop("DOCCONVERT_WORKSPACE", None)
            
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_import_valid_docx(self):
        """Test importing a valid Word .docx file into Markdown."""
        docx_path = os.path.join(self.workspace_dir, "report.docx")
        doc = docx.Document()
        doc.add_heading("Tiêu đề Báo Cáo", level=1)
        doc.add_paragraph("Nội dung thử nghiệm bóc tách Word sang Markdown.")
        doc.save(docx_path)

        res = handle_import_document_to_markdown(docx_path, index=self.index)
        self.assertEqual(res.get("status"), "success")
        self.assertIn("document_id", res)
        
        expected_md = os.path.join(self.workspace_dir, "report.md")
        self.assertTrue(os.path.exists(expected_md))
        
        # Check read_document retrieves it
        read_res = handle_read_document(res["document_id"], index=self.index)
        self.assertNotIn("error", read_res)
        self.assertIn("Tiêu đề Báo Cáo", read_res["content"])

    def test_import_valid_xlsx(self):
        """Test importing an Excel .xlsx sheet into Markdown table."""
        xlsx_path = os.path.join(self.workspace_dir, "data.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Tên Sản Phẩm", "Số Lượng", "Đơn Giá"])
        ws.append(["Bàn phím cơ", 10, 1500000])
        wb.save(xlsx_path)

        res = handle_import_document_to_markdown(xlsx_path, index=self.index)
        self.assertEqual(res.get("status"), "success")
        
        expected_md = os.path.join(self.workspace_dir, "data.md")
        self.assertTrue(os.path.exists(expected_md))

    def test_security_path_outside_workspace(self):
        """Test that importing files outside workspace is rejected (G4)."""
        outside_file = os.path.join(self.tmp_dir, "outside.docx")
        doc = docx.Document()
        doc.add_paragraph("Secret outside content")
        doc.save(outside_file)

        res = handle_import_document_to_markdown(outside_file, index=self.index)
        self.assertEqual(res.get("error"), "PATH_OUTSIDE_WORKSPACE")

    def test_security_traversal_relative_path(self):
        """Test path traversal with ../ is blocked."""
        res = handle_import_document_to_markdown("../../secret.docx", index=self.index)
        self.assertEqual(res.get("error"), "PATH_OUTSIDE_WORKSPACE")

    def test_unsupported_format(self):
        """Test that unsupported formats are rejected."""
        txt_path = os.path.join(self.workspace_dir, "script.py")
        with open(txt_path, "w") as f:
            f.write("print('hello')")

        res = handle_import_document_to_markdown(txt_path, index=self.index)
        self.assertEqual(res.get("error"), "UNSUPPORTED_FORMAT")

    def test_pdf_page_limit_exceeded(self):
        """Test that PDF > 30 pages is rejected with PDF_PAGE_LIMIT_EXCEEDED."""
        pdf_path = os.path.join(self.workspace_dir, "large_31p.pdf")
        doc = fitz.open()
        for i in range(31):
            p = doc.new_page()
            p.insert_text((50, 50), f"Page {i+1} text content")
        doc.save(pdf_path)
        doc.close()

        res = handle_import_document_to_markdown(pdf_path, index=self.index)
        self.assertEqual(res.get("error"), "PDF_PAGE_LIMIT_EXCEEDED")

    def test_pdf_zero_pages(self):
        """Test that PDF with 0 pages is rejected with EMPTY_DOCUMENT."""
        pdf_path = os.path.join(self.workspace_dir, "mock_0p.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4\n%%EOF")

        with patch("fitz.open") as mock_fitz:
            mock_doc = MagicMock()
            mock_doc.__len__.return_value = 0
            mock_doc.is_encrypted = False
            mock_doc.needs_pass = False
            mock_fitz.return_value.__enter__.return_value = mock_doc

            res = handle_import_document_to_markdown(pdf_path, index=self.index)
            self.assertEqual(res.get("error"), "EMPTY_DOCUMENT")

    def test_scanned_pdf_empty_content(self):
        """Test that an image-only scanned PDF with no text returns EMPTY_CONTENT."""
        pdf_path = os.path.join(self.workspace_dir, "scan_only.pdf")
        img_path = os.path.join(self.workspace_dir, "temp_img.png")
        img = Image.new("RGB", (200, 200), color=(255, 255, 255))
        img.save(img_path)

        doc = fitz.open()
        p = doc.new_page(width=200, height=200)
        p.insert_image(fitz.Rect(0, 0, 200, 200), filename=img_path)
        doc.save(pdf_path)
        doc.close()
        os.remove(img_path)

        res = handle_import_document_to_markdown(pdf_path, index=self.index)
        self.assertEqual(res.get("error"), "EMPTY_CONTENT")
        self.assertIn("OCR", res.get("message", ""))

    def test_zip_bomb_detection(self):
        """Test zip bomb protection for oversized uncompressed XML."""
        fake_docx = os.path.join(self.workspace_dir, "bomb.docx")
        import zipfile
        with zipfile.ZipFile(fake_docx, 'w', compression=zipfile.ZIP_DEFLATED) as z:
            # Write a mock huge entry > 50MB
            huge_data = b"0" * (52 * 1024 * 1024)
            z.writestr("word/document.xml", huge_data)

        res = handle_import_document_to_markdown(fake_docx, index=self.index)
        self.assertEqual(res.get("error"), "ZIP_BOMB_DETECTED")

    def test_corrupted_office_archive(self):
        """Test that corrupted non-zip file with .docx extension returns IMPORT_FAILED."""
        bad_docx = os.path.join(self.workspace_dir, "corrupt.docx")
        with open(bad_docx, "wb") as f:
            f.write(b"NOT_A_VALID_ZIP_HEADER_DATA")

        res = handle_import_document_to_markdown(bad_docx, index=self.index)
        self.assertEqual(res.get("error"), "IMPORT_FAILED")

    def test_filename_stem_sanitization(self):
        """Test filename sanitization against Windows device names and illegal characters."""
        self.assertEqual(_sanitize_filename_stem("NUL.tar.gz", fallback_stem="fallback"), "fallback")
        self.assertEqual(_sanitize_filename_stem("CON", fallback_stem="fallback"), "fallback")
        self.assertEqual(_sanitize_filename_stem("AUX.docx", fallback_stem="fallback"), "fallback")
        self.assertEqual(_sanitize_filename_stem("x.md:hidden_stream", fallback_stem="fallback"), "x")
        self.assertEqual(_sanitize_filename_stem("valid_name.docx", fallback_stem="fallback"), "valid_name")

    def test_error_message_sanitization(self):
        """Test that system drive paths and UNC paths are stripped from error messages."""
        msg_win = "Error loading file C:\\Users\\Admin\\Secret\\doc.docx failed"
        clean = _sanitize_error_message(msg_win)
        self.assertNotIn("C:\\Users\\Admin", clean)
        self.assertIn("<path>", clean)

        msg_unc = "Error opening \\\\server\\share\\secret.pdf"
        clean_unc = _sanitize_error_message(msg_unc)
        self.assertNotIn("\\\\server\\share", clean_unc)
        self.assertIn("<path>", clean_unc)

    def test_session_only_asset_containment_in_export_assets(self):
        """Test that export_assets with session_only=True refuses to copy external disk images."""
        mgr = MediaAssetManager()
        session_id = "test_containment_sess"
        mgr.set_active_session(session_id)

        # 1. Image outside session cache
        secret_img = os.path.join(self.tmp_dir, "secret.png")
        Image.new("RGB", (10, 10), color=(0, 0, 0)).save(secret_img)

        # 2. Image inside session cache
        sess_dir = mgr.get_session_dir(session_id)
        valid_img = os.path.join(sess_dir, "valid.png")
        Image.new("RGB", (10, 10), color=(255, 255, 255)).save(valid_img)

        md_input = f"![Valid](@media/valid.png) and ![Secret]({secret_img})"
        target_md = os.path.join(self.workspace_dir, "output.md")

        # Run with session_only=True
        exported = mgr.export_assets(md_input, target_md, session_id=session_id, session_only=True)

        assets_folder = os.path.join(self.workspace_dir, "output_assets")
        self.assertTrue(os.path.exists(os.path.join(assets_folder, "valid.png")))
        # Secret image must NOT be copied to assets folder
        self.assertFalse(os.path.exists(os.path.join(assets_folder, "secret.png")))
        self.assertIn(secret_img, exported)  # Uncopied external link remains untouched

        mgr.clear_session(session_id)


    def test_lifecycle_import_patch_rollback(self):
        """Test full document lifecycle: import -> patch_document_content -> rollback_document."""
        docx_path = os.path.join(self.workspace_dir, "lifecycle.docx")
        doc = docx.Document()
        doc.add_paragraph("Original text line 1.")
        doc.save(docx_path)

        # 1. Import
        imp_res = handle_import_document_to_markdown(docx_path, index=self.index)
        self.assertEqual(imp_res.get("status"), "success")
        doc_id = imp_res["document_id"]

        # 2. Patch
        patch_res = handle_patch_document_content(
            document_id=doc_id,
            target_content="Original text line 1.",
            replacement_content="Patched text line 1 by AI.",
            index=self.index
        )
        self.assertEqual(patch_res.get("status"), "success")

        # Verify patch applied
        read_after_patch = handle_read_document(doc_id, index=self.index)
        self.assertIn("Patched text line 1 by AI.", read_after_patch["content"])

        # 3. Rollback to version 0 (initial baseline)
        roll_res = handle_rollback_document(doc_id, target_version=0, index=self.index)
        self.assertEqual(roll_res.get("status"), "success")

        # Verify rollback restored original baseline
        read_after_roll = handle_read_document(doc_id, index=self.index)
        self.assertIn("Original text line 1.", read_after_roll["content"])
        self.assertNotIn("Patched text line 1 by AI.", read_after_roll["content"])


    def test_session_only_false_regression(self):
        """Regression test: when session_only=False (GUI default), external disk images are copied as before."""
        mgr = MediaAssetManager()
        session_id = "test_regression_sess"
        mgr.set_active_session(session_id)

        external_img = os.path.join(self.tmp_dir, "gui_ext.png")
        Image.new("RGB", (10, 10), color=(100, 100, 100)).save(external_img)

        md_input = f"![GUI External]({external_img})"
        target_md = os.path.join(self.workspace_dir, "gui_out.md")

        # Run with session_only=False (default)
        exported = mgr.export_assets(md_input, target_md, session_id=session_id, session_only=False)

        assets_folder = os.path.join(self.workspace_dir, "gui_out_assets")
        self.assertTrue(os.path.exists(os.path.join(assets_folder, "gui_ext.png")))
        self.assertIn("./gui_out_assets/gui_ext.png", exported)

        mgr.clear_session(session_id)

    def test_output_too_large_rejection(self):
        """Test that Markdown output exceeding 5MB is rejected with OUTPUT_TOO_LARGE."""
        csv_path = os.path.join(self.workspace_dir, "large.csv")
        # Write 6MB CSV
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("col1,col2,col3\n")
            for _ in range(120000):
                f.write("data_sample_value_1,data_sample_value_2,data_sample_value_3\n")

        with patch("src.services.file_loader.load_document") as mock_load:
            from src.services.file_loader import LoadResult
            mock_load.return_value = LoadResult(
                success=True,
                content="A" * (6 * 1024 * 1024),  # 6MB string
                path=csv_path
            )
            res = handle_import_document_to_markdown(csv_path, index=self.index)
            self.assertEqual(res.get("error"), "OUTPUT_TOO_LARGE")

    def test_encrypted_pdf_rejection(self):
        """Test that encrypted/password-protected PDF is rejected with ENCRYPTED_DOCUMENT."""
        pdf_path = os.path.join(self.workspace_dir, "encrypted.pdf")
        with open(pdf_path, "wb") as f:
            f.write(b"%PDF-1.4\n%%EOF")

        with patch("fitz.open") as mock_fitz:
            mock_doc = MagicMock()
            mock_doc.__len__.return_value = 5
            mock_doc.is_encrypted = True
            mock_doc.needs_pass = True
            mock_fitz.return_value.__enter__.return_value = mock_doc

            res = handle_import_document_to_markdown(pdf_path, index=self.index)
            self.assertEqual(res.get("error"), "ENCRYPTED_DOCUMENT")

    def test_version_baseline_warning_does_not_delete_file(self):
        """Test that failure in ensure_baseline_version returns a warning and does NOT delete the indexed file."""
        docx_path = os.path.join(self.workspace_dir, "warning_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("Content for baseline failure test.")
        doc.save(docx_path)

        with patch.object(self.index, "ensure_baseline_version", side_effect=RuntimeError("Simulated DB lock")):
            res = handle_import_document_to_markdown(docx_path, index=self.index)
            self.assertEqual(res.get("status"), "success")
            self.assertIn("warning", res)
            self.assertTrue(os.path.exists(os.path.join(self.workspace_dir, "warning_doc.md")))

    def test_json_rpc_subprocess_execution(self):
        """Test invoking import_document_to_markdown via actual JSON-RPC stdio subprocess."""
        import subprocess

        docx_path = os.path.join(self.workspace_dir, "rpc_test.docx")
        doc = docx.Document()
        doc.add_paragraph("Testing MCP stdio JSON-RPC pipeline.")
        doc.save(docx_path)

        req = {
            "jsonrpc": "2.0",
            "id": 1,
            "method": "tools/call",
            "params": {
                "name": "import_document_to_markdown",
                "arguments": {
                    "source_path": docx_path
                }
            }
        }
        body = json.dumps(req)
        raw_msg = f"Content-Length: {len(body.encode('utf-8'))}\r\n\r\n{body}"

        env = dict(os.environ)
        env["DOCCONVERT_WORKSPACE"] = self.workspace_dir
        env["DOCCONVERT_INDEX_PATH"] = self.db_path

        proc = subprocess.Popen(
            [sys.executable, "-m", "src.mcp.server"],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=False,
            env=env,
            cwd=os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        )

        stdout_data, stderr_data = proc.communicate(input=raw_msg.encode("utf-8"), timeout=10)
        
        # Verify stdout only contains valid JSON-RPC Content-Length messages
        stdout_text = stdout_data.decode("utf-8", errors="replace")
        self.assertIn("jsonrpc", stdout_text)
        
        # Parse returned JSON-RPC payload directly
        json_body = json.loads(stdout_text.strip())
        tool_res = json_body["result"]["content"][0]["text"]
        parsed_res = json.loads(tool_res)
        
        self.assertEqual(parsed_res.get("status"), "success")
        self.assertIn("document_id", parsed_res)



if __name__ == "__main__":
    unittest.main()

